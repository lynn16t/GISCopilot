"""
SpatialAnalysisAgent_KnowledgeManager.py

Manages project-level knowledge: user-written semantic notes and uploaded
reference documents.  Knowledge is stored alongside the QGIS project
(or in the plugin workspace when no project is open).

Storage layout
--------------
<knowledge_root>/
    notes.md              – free-form user notes (data dictionaries, rules …)
    docs/                 – uploaded reference documents (PDF, DOCX, XLSX, TXT …)
    docs_text/            – extracted plain-text mirrors of each uploaded doc

Public API used by the rest of the plugin
-----------------------------------------
    get_notes()           -> str
    save_notes(text)
    add_document(src_path) -> str          # returns display name
    remove_document(filename)
    list_documents()      -> list[dict]    # [{filename, size_kb, added}]
    get_relevant_knowledge(layer_names, max_chars) -> str
"""

import os
import shutil
import json
import datetime
from pathlib import Path


# ---------------------------------------------------------------------------
# Text extraction helpers  (pure Python, no LLM)
# ---------------------------------------------------------------------------

def _extract_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _extract_pdf(path: str) -> str:
    """Extract text from PDF using PyMuPDF (fitz) or fallback to pdfplumber."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(path)
        pages = [page.get_text() for page in doc]
        doc.close()
        return "\n\n".join(pages)
    except ImportError:
        pass
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n\n".join(p.extract_text() or "" for p in pdf.pages)
    except ImportError:
        return f"[PDF text extraction unavailable – install PyMuPDF or pdfplumber]"


def _extract_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except ImportError:
        return "[DOCX extraction unavailable – install python-docx]"


def _extract_xlsx(path: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"--- Sheet: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                line = "\t".join(str(c) if c is not None else "" for c in row)
                if line.strip():
                    lines.append(line)
        wb.close()
        return "\n".join(lines)
    except ImportError:
        return "[XLSX extraction unavailable – install openpyxl]"


def _extract_csv(path: str) -> str:
    return _extract_txt(path)


_EXTRACTORS = {
    ".txt": _extract_txt,
    ".md":  _extract_txt,
    ".csv": _extract_csv,
    ".tsv": _extract_csv,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".xls":  _extract_xlsx,
    ".json": _extract_txt,
    ".toml": _extract_txt,
}


# ---------------------------------------------------------------------------
# ProjectKnowledgeManager
# ---------------------------------------------------------------------------

class ProjectKnowledgeManager:
    """Singleton-style manager bound to a knowledge root directory."""

    def __init__(self, knowledge_root: str | None = None):
        self._root = None
        if knowledge_root:
            self.set_root(knowledge_root)

    # ------------------------------------------------------------------
    # Root directory management
    # ------------------------------------------------------------------

    def set_root(self, knowledge_root: str):
        """Set (or change) the knowledge root and ensure directory structure."""
        self._root = Path(knowledge_root)
        self._notes_path = self._root / "notes.md"
        self._docs_dir = self._root / "docs"
        self._text_dir = self._root / "docs_text"
        self._meta_path = self._root / "docs_meta.json"

        self._root.mkdir(parents=True, exist_ok=True)
        self._docs_dir.mkdir(exist_ok=True)
        self._text_dir.mkdir(exist_ok=True)

        # Create empty notes file if it doesn't exist
        if not self._notes_path.exists():
            self._notes_path.write_text("", encoding="utf-8")

        # Create empty metadata file if it doesn't exist
        if not self._meta_path.exists():
            self._meta_path.write_text("[]", encoding="utf-8")

    @property
    def root(self) -> str | None:
        return str(self._root) if self._root else None

    @property
    def is_ready(self) -> bool:
        return self._root is not None and self._root.exists()

    # ------------------------------------------------------------------
    # Notes (free-form text editor content)
    # ------------------------------------------------------------------

    def get_notes(self) -> str:
        if not self.is_ready:
            return ""
        try:
            return self._notes_path.read_text(encoding="utf-8")
        except Exception:
            return ""

    def save_notes(self, text: str):
        if not self.is_ready:
            return
        self._notes_path.write_text(text, encoding="utf-8")

    # ------------------------------------------------------------------
    # Document management
    # ------------------------------------------------------------------

    def _load_meta(self) -> list:
        try:
            return json.loads(self._meta_path.read_text(encoding="utf-8"))
        except Exception:
            return []

    def _save_meta(self, meta: list):
        self._meta_path.write_text(
            json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    def add_document(self, src_path: str) -> str:
        """Copy a document into the knowledge store and extract its text.

        Returns the display filename.
        """
        if not self.is_ready:
            raise RuntimeError("Knowledge root not set")

        src = Path(src_path)
        dest = self._docs_dir / src.name

        # Handle name collision
        counter = 1
        while dest.exists():
            dest = self._docs_dir / f"{src.stem}_{counter}{src.suffix}"
            counter += 1

        shutil.copy2(str(src), str(dest))

        # Extract text
        ext = dest.suffix.lower()
        extractor = _EXTRACTORS.get(ext)
        if extractor:
            try:
                text = extractor(str(dest))
            except Exception as e:
                text = f"[Extraction failed: {e}]"
        else:
            text = f"[Unsupported format: {ext}]"

        text_file = self._text_dir / (dest.stem + ".txt")
        text_file.write_text(text, encoding="utf-8")

        # Update metadata
        meta = self._load_meta()
        meta.append({
            "filename": dest.name,
            "original_path": str(src),
            "size_kb": round(dest.stat().st_size / 1024, 1),
            "added": datetime.datetime.now().isoformat(timespec="seconds"),
            "text_file": text_file.name,
        })
        self._save_meta(meta)

        return dest.name

    def remove_document(self, filename: str):
        """Remove a document and its extracted text from the store."""
        if not self.is_ready:
            return

        meta = self._load_meta()
        entry = next((m for m in meta if m["filename"] == filename), None)
        if not entry:
            return

        # Delete files
        doc_path = self._docs_dir / filename
        if doc_path.exists():
            doc_path.unlink()

        text_file = self._text_dir / entry.get("text_file", "")
        if text_file.exists():
            text_file.unlink()

        # Update metadata
        meta = [m for m in meta if m["filename"] != filename]
        self._save_meta(meta)

    def list_documents(self) -> list:
        """Return list of dicts with filename, size_kb, added."""
        if not self.is_ready:
            return []
        return self._load_meta()

    # ------------------------------------------------------------------
    # Knowledge retrieval for LLM context injection
    # ------------------------------------------------------------------

    def get_relevant_knowledge(
        self,
        layer_names: list[str] | None = None,
        query: str = "",
        max_chars: int = 8000,
    ) -> str:
        """Build a knowledge string to inject into LLM context.

        Strategy (Option A – raw text, effect-first):
        1. Always include the full notes.md (user-curated, high value)
        2. For each uploaded document, include extracted text up to budget

        If layer_names or query are provided, prioritize documents whose
        text mentions those terms (simple keyword relevance).
        """
        if not self.is_ready:
            return ""

        parts = []
        budget = max_chars

        # --- Notes (always first, always included) ---
        notes = self.get_notes().strip()
        if notes:
            section = f"=== Project Notes ===\n{notes}"
            parts.append(section)
            budget -= len(section)

        if budget <= 0:
            return "\n\n".join(parts)

        # --- Uploaded documents ---
        meta = self._load_meta()
        if not meta:
            return "\n\n".join(parts)

        # Build search terms for relevance scoring
        search_terms = []
        if layer_names:
            search_terms.extend([n.lower() for n in layer_names])
        if query:
            search_terms.extend(query.lower().split())

        # Score and sort documents by relevance
        scored_docs = []
        for entry in meta:
            text_path = self._text_dir / entry.get("text_file", "")
            if not text_path.exists():
                continue
            try:
                text = text_path.read_text(encoding="utf-8")
            except Exception:
                continue

            if not text.strip():
                continue

            # Simple keyword relevance score
            score = 0
            text_lower = text.lower()
            for term in search_terms:
                if term in text_lower:
                    score += text_lower.count(term)

            # If no search terms, all docs are equally relevant
            if not search_terms:
                score = 1

            scored_docs.append((score, entry["filename"], text))

        # Sort by relevance (highest first)
        scored_docs.sort(key=lambda x: x[0], reverse=True)

        # Pack documents into budget
        for score, filename, text in scored_docs:
            if budget <= 200:  # reserve a bit for formatting
                break

            # Truncate individual doc if needed
            available = min(len(text), budget - 50)
            if available <= 0:
                break

            snippet = text[:available]
            if len(text) > available:
                snippet += "\n... [truncated]"

            section = f"=== Document: {filename} ===\n{snippet}"
            parts.append(section)
            budget -= len(section)

        return "\n\n".join(parts)

    # ------------------------------------------------------------------
    # Utility
    # ------------------------------------------------------------------

    def get_knowledge_root_for_project(
        self, workspace_dir: str, qgis_project_path: str | None = None
    ) -> str:
        """Determine the best knowledge root directory.

        If a QGIS project file (.qgz/.qgs) is open, store knowledge
        next to it.  Otherwise fall back to the plugin workspace.
        """
        if qgis_project_path:
            project_dir = os.path.dirname(qgis_project_path)
            if os.path.isdir(project_dir):
                return os.path.join(project_dir, "gis_knowledge")

        return os.path.join(workspace_dir, "gis_knowledge")
